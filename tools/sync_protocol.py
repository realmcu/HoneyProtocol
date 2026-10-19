#!/usr/bin/env python3
"""Generate the normative protocol pages from the reviewed Excel workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/assets/protocol/cavo-realtek-watch-protocol.xlsx"
OUTPUT_DIR = ROOT / "docs/protocol"
MANIFEST = OUTPUT_DIR / "source.json"

HEADER_FILLS = {"FFD9E1F2", "FFBDD7EE"}
HEADER_TITLES = {
    "字段",
    "长度",
    "取值",
    "说明",
    "操作码",
    "方向",
    "操作数据",
    "Key",
    "KeyID",
    "Key 范围",
    "名称",
    "分组",
    "场景",
    "具体的命令",
    "数据说明",
    "Value",
    "编号",
    "事项",
    "想要",
    "类型",
    "index",
    "筛选条件",
    "偏移",
    "改动",
    "处理",
    "错误码",
    "状态码",
    "Data type",
}
HEX_PAIR = re.compile(r"(?<![0-9A-Fa-f])[0-9A-Fa-f]{2}(?![0-9A-Fa-f])")
PACKET = re.compile(r"^(?:[0-9A-Fa-f]{2})(?:\s+[0-9A-Fa-f]{2}){4,}$")
PACKET_WITH_GAP = re.compile(
    r"^[0-9A-Fa-f]{2}(?:\s+[0-9A-Fa-f]{2})*(?:\s+(?:…|\.\.\.)(?:\s+[0-9A-Fa-f]{2})*)+$"
)


@dataclass(frozen=True)
class Document:
    sheet: str
    filename: str
    title: str
    command: str | None = None


DOCUMENTS = (
    Document("设置模块(CMD 02)", "0x02-settings.rst", "0x02 设置命令", "02"),
    Document("绑定模块(CMD 03)", "0x03-bind.rst", "0x03 绑定命令", "03"),
    Document("通知模块(CMD 04)", "0x04-notification.rst", "0x04 通知命令", "04"),
    Document("数据模块(CMD 05)", "0x05-data.rst", "0x05 数据命令", "05"),
    Document("点测模块(CMD 11)", "0x11-measurement.rst", "0x11 点测命令", "11"),
    Document("模拟模块(CMD 12)", "0x12-simulator.rst", "0x12 模拟数据命令", "12"),
    Document("音频模块(CMD 13)", "0x13-audio.rst", "0x13 音频命令", "13"),
    Document("日志模块(CMD 0A)", "0x0a-log.rst", "0x0A 日志命令", "0A"),
    Document("天气类型", "weather-types.rst", "天气类型"),
    Document("语音编码", "language-codes.rst", "语言编码"),
)

# Commands absent from the workbook; their pages in OUTPUT_DIR are maintained by hand.
MANUAL_DOCUMENTS = (
    ("0x01-firmware-upgrade.rst", "0x01 固件升级命令", "01"),
    ("0x06-factory-test.rst", "0x06 工厂测试命令", "06"),
    ("0x08-dump-stack.rst", "0x08 Dump Stack 命令", "08"),
    ("0x09-flash-read.rst", "0x09 测试 Flash 读取命令", "09"),
    ("0x0c-ble-conn-param.rst", "0x0C BLE 连接参数命令", "0C"),
    ("0x0d-wifi-provision.rst", "0x0D WiFi 配网命令", "0D"),
    ("0x10-file-transfer.rst", "0x10 文件传输命令", "10"),
)

RST_HEADING_CHARS = ("=", "-", "~", "^")
SECTION_NAMES = {
    "Key 列表": "消息定义",
    "Key（消息类型）": "消息定义",
    "表格说明": "具体命令",
}
CHINESE_SECTION_NUMBERS = {
    "一": "1",
    "二": "2",
    "三": "3",
    "四": "4",
    "五": "5",
    "六": "6",
}
RELATED_DOCUMENTS = {
    "ws-protocol-notify.md": ("通知命令", "0x04-notification"),
    "ws-protocol-data.md": ("数据命令", "0x05-data"),
    "ws-protocol-measure.md": ("点测命令", "0x11-measurement"),
}
SHEET_DOCUMENTS = {
    "天气类型": ("天气类型", "weather-types"),
    "语音编码": ("语言编码", "language-codes"),
    "语言编码": ("语言编码", "language-codes"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def fill_rgb(cell: Cell) -> str | None:
    fill = cell.fill
    if not fill.fill_type:
        return None
    color = fill.start_color
    if isinstance(color.rgb, str):
        return color.rgb
    return None


def is_table_header(ws, row_number: int) -> bool:
    marked = []
    nonempty = []
    for column in range(1, ws.max_column + 1):
        cell = ws.cell(row_number, column)
        if cell_text(cell.value):
            nonempty.append(cell)
            if fill_rgb(cell) in HEADER_FILLS:
                marked.append(cell)
    if len(marked) < 2 or len(nonempty) < 2:
        marked_titles = sum(cell_text(cell.value) in HEADER_TITLES for cell in nonempty)
        return marked_titles >= 2 and all(
            len(cell_text(cell.value)) <= 32 for cell in nonempty
        )
    return all(len(cell_text(cell.value)) <= 32 for cell in nonempty)


def is_blank_header_spacer(ws, row_number: int) -> bool:
    return not any(cell_text(ws.cell(row_number, column).value) for column in range(1, ws.max_column + 1)) and sum(
        fill_rgb(ws.cell(row_number, column)) in HEADER_FILLS
        for column in range(1, ws.max_column + 1)
    ) >= 2


def read_grid(ws) -> list[list[str]]:
    grid = [
        [cell_text(ws.cell(row, column).value) for column in range(1, ws.max_column + 1)]
        for row in range(1, ws.max_row + 1)
    ]
    for merged in ws.merged_cells.ranges:
        if merged.min_col != merged.max_col or merged.min_row == merged.max_row:
            continue
        value = grid[merged.min_row - 1][merged.min_col - 1]
        if not value:
            continue
        for row in range(merged.min_row + 1, merged.max_row + 1):
            grid[row - 1][merged.min_col - 1] = value
    return grid


def table_end(ws, start: int, grid: list[list[str]]) -> int:
    row = start + 1
    while row <= len(grid):
        if is_table_header(ws, row):
            break
        if is_blank_header_spacer(ws, row):
            row += 1
            continue
        if sum(bool(value) for value in grid[row - 1]) < 2:
            break
        row += 1
    return row - 1


def inline_rst(value: str) -> str:
    value = " ".join(value.splitlines())
    value = re.sub(r"`([^`]+)`", r"``\1``", value)
    parts = re.split(r"(\*\*.*?\*\*)", value)
    value = "".join(part if part.startswith("**") and part.endswith("**") else part.replace("*", r"\*") for part in parts)
    if value and (PACKET.fullmatch(value) or PACKET_WITH_GAP.fullmatch(value)):
        value = f"``{value}``"
    for reference, (label, target) in RELATED_DOCUMENTS.items():
        value = value.replace(reference, f":doc:`{label} <{target}>`")
    for sheet, (label, target) in SHEET_DOCUMENTS.items():
        value = re.sub(
            rf'[“"]{sheet}[”"](?:工作表)?',
            lambda match: f":doc:`{label} <{target}>`",
            value,
        )
    value = value.replace("HL_Private_Protocol.md", "旧版协议（原文档未随仓库提供）")
    value = value.replace("BLE_PROTOCOL_SPEC.html", "旧版协议（原文档未随仓库提供）")
    return value


def heading(text: str, level: int = 1) -> str:
    char = RST_HEADING_CHARS[min(level, len(RST_HEADING_CHARS)) - 1]
    width = sum(2 if unicodedata.east_asian_width(letter) in "WF" else 1 for letter in text)
    return f"{text}\n{char * width}"


def render_text_lines(lines: list[str]) -> str:
    output = []
    previous_kind = None
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            if output and output[-1]:
                output.append("")
            previous_kind = None
            continue
        bullet = line.startswith("·") or line.startswith("-")
        if bullet:
            line = inline_rst(line.lstrip("·-").strip())
            line = f"* {line}"
        else:
            line = inline_rst(line)
        if re.match(r"^\d+\.\s+", line):
            kind = "ordered-list"
        elif re.match(r"^(?:\*|-)\s+", line):
            kind = "unordered-list"
        else:
            kind = "text"
        if output and output[-1] and kind != previous_kind:
            output.append("")
        output.append(line)
        previous_kind = kind
    return "\n".join(output).strip()


def render_table(ws, start: int, end: int, grid: list[list[str]]) -> str:
    rows = [row for row in grid[start - 1 : end] if any(row)]
    used_columns = [
        column
        for column in range(len(grid[0]))
        if any(row[column] for row in rows)
    ]
    output = [
        ".. list-table::",
        "   :header-rows: 1",
        "   :widths: auto",
        f"   :class: protocol-table columns-{len(used_columns)}",
        "",
    ]
    for row in rows:
        values = [inline_rst(row[column]) or "—" for column in used_columns]
        output.append(f"   * - {values[0]}")
        output.extend(f"     - {value}" for value in values[1:])
    return "\n".join(output)


def narrative_heading(text: str) -> int | None:
    if "\n" in text or len(text) > 80:
        return None
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return 2
    if re.match(r"^\d+\.\s+(?:通用说明|具体命令)$", text):
        return 2
    if text == "具体命令":
        return 2
    match = re.match(r"^(\d+(?:\.\d+)+)\s+\S", text)
    return min(4, 2 + match.group(1).count(".")) if match else None


def section_title(text: str) -> str | None:
    if not narrative_heading(text):
        return None
    name = re.sub(r"^(?:[一二三四五六七八九十]+、|\d+(?:\.\d+)*\.?\s+)", "", text, count=1)
    if name == "通用说明":
        return ""
    if name.startswith("命令格式"):
        return "命令格式"
    if name.startswith("操作码"):
        return "操作码"
    if name == "旧版协议差异" or name.startswith("与老协议（"):
        return "旧版协议差异"
    return SECTION_NAMES.get(name, name)


def render_example(text: str) -> str:
    text = re.sub(
        r"(Value|操作数据)(?=[0-9A-Fa-f]{2}(?:\s+[0-9A-Fa-f]{2})*\s*\|)",
        r"\1\n",
        text,
    )
    label, *lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) >= 2 and lines[0].startswith("L2 Header"):
        fields = [part.strip() for part in lines[0].split("|")]
        values = [part.strip() for part in lines[1].split("|")]
        if len(fields) == len(values):
            output = [
                f".. list-table:: {inline_rst(label)}",
                "   :header-rows: 1",
                "   :widths: auto",
                "   :class: packet-table",
                "",
                f"   * - {inline_rst(fields[0])}",
                *[f"     - {inline_rst(field)}" for field in fields[1:]],
                f"   * - ``{values[0]}``",
                *[f"     - ``{value}``" for value in values[1:]],
            ]
            if len(lines) > 2:
                output.extend(["", render_text_lines(lines[2:])])
            return "\n".join(output)
    return f"**{inline_rst(label)}**\n\n.. code-block:: text\n\n" + "\n".join(
        f"   {line}" for line in lines
    )


def render_flow(text: str) -> str:
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "手机 手表")
    end = start + 1
    while end < len(lines) and not re.match(r"^\d+\.\s+", lines[end]):
        end += 1
    parts = []
    if start:
        parts.append(render_text_lines(lines[:start]))
    parts.append(
        ".. code-block:: text\n\n"
        + "\n".join(f"   {line}" for line in lines[start:end])
    )
    if end < len(lines):
        parts.append(render_text_lines(lines[end:]))
    return "\n\n".join(parts)


def render_narrative(values: Iterable[str]) -> str:
    parts = []
    prose = []
    for value in values:
        if "灰色：L2 Header" in value:
            continue
        value = value.replace(
            '，可直接对照抓包。按字段分颜色：',
            '，可直接对照抓包。',
        )
        if value.startswith("示例：") or "手机 手表\n" in value:
            if prose:
                parts.append(render_text_lines(prose))
                prose = []
            parts.append(render_example(value) if value.startswith("示例：") else render_flow(value))
        else:
            prose.extend(value.splitlines())
    if prose:
        parts.append(render_text_lines(prose))
    return "\n\n".join(part for part in parts if part)


def render_sheet(document: Document, ws) -> tuple[str, str | None]:
    grid = read_grid(ws)
    status = None
    if len(grid) > 1 and grid[1][0].startswith("状态："):
        status = grid[1][0].removeprefix("状态：").strip()

    blocks = []
    sections = {}
    last_section = None
    last_label = None
    row = 1 if document.command is None else 2
    while row <= len(grid):
        if row == 2 and status:
            row += 1
            continue
        if not any(grid[row - 1]):
            row += 1
            continue
        if is_table_header(ws, row):
            end = table_end(ws, row, grid)
            blocks.append(render_table(ws, row, end, grid))
            row = end + 1
            continue
        current = [value for value in grid[row - 1] if value]
        title = section_title(current[0]) if len(current) == 1 else None
        if title is not None:
            number = re.match(r"^(\d+(?:\.\d+)*)(?:\.\s+|\s+)", current[0])
            chinese_number = re.match(r"^([一二三四五六])、", current[0])
            section_number = (
                number.group(1)
                if number
                else CHINESE_SECTION_NUMBERS[chinese_number.group(1)] if chinese_number else None
            )
            if title and title != last_section:
                if section_number:
                    label = f"cmd-{document.command.lower()}-section-{section_number.replace('.', '-')}"
                    sections[section_number] = (title, label)
                    blocks.append(f".. _{label}:")
                    last_label = label
                else:
                    last_label = None
                blocks.append(heading(title, 2))
                last_section = title
            elif title and section_number and last_label:
                sections[section_number] = (title, last_label)
            if title == "命令格式":
                blocks.append("以下为 L2 层命令格式，不含 L1 帧头。")
            if title == "旧版协议差异":
                blocks.append("源工作簿引用的旧版协议原文未随仓库提供，以下差异待对照原文复核。")
            row += 1
            continue
        values = []
        while row <= len(grid) and any(grid[row - 1]) and not is_table_header(ws, row):
            current = [value for value in grid[row - 1] if value]
            if len(current) == 1 and section_title(current[0]) is not None:
                break
            values.append(" ".join(current))
            row += 1
        block = render_narrative(values)
        if block:
            blocks.append(block)

    header = [
        ":tocdepth: 2" if document.command else ":orphan:",
        heading(document.title),
        ".. container:: doc-meta\n\n"
        f"   来源：:download:`协议工作簿 <../assets/protocol/{SOURCE.name}>` · "
        f"``{document.sheet}`` 工作表（自动同步）",
    ]
    if status:
        label, _, summary = status.partition("。")
        header.append(f".. container:: doc-status\n\n   **{inline_rst(label)}**　{inline_rst(summary)}")
    page = "\n\n".join(header + blocks).rstrip() + "\n"
    for number, (title, label) in sorted(sections.items(), key=lambda item: -len(item[0])):
        page = re.sub(
            rf"(?<=(?:见|按|同) ){re.escape(number)}(?![\d.])",
            f":ref:`{title} <{label}>`",
            page,
        )
        page = re.sub(
            rf"见第\s*{re.escape(number)}\s*节",
            f"见 :ref:`{title} <{label}>`",
            page,
        )
        page = re.sub(
            rf"第\s*{re.escape(number)}\s*节",
            f":ref:`{title} <{label}>`",
            page,
        )
        for chinese, value in CHINESE_SECTION_NUMBERS.items():
            if number == value:
                page = re.sub(
                    rf'见[“"]{chinese}、[^”"]+[”"]',
                    f"见 :ref:`{title} <{label}>`",
                    page,
                )
    return page, status


def render_index(summaries: list[dict], missing_sheets: dict[str, list[str]]) -> str:
    commands = [
        (item["command"], item["title"], item["filename"], item["status"], "工作簿")
        for item in summaries
        if item["command"] is not None
    ]
    commands.extend(
        (command, title, filename, None, "手工维护") for filename, title, command in MANUAL_DOCUMENTS
    )
    commands.sort(key=lambda entry: int(entry[0], 16))
    rows = [
        ".. list-table:: Command 列表",
        "   :header-rows: 1",
        "   :widths: auto",
        "",
        "   * - Command ID",
        "     - 命令",
        "     - 规范状态",
        "     - 来源",
    ]
    for command, title, filename, status, source in commands:
        name = title.removeprefix(f"0x{command} ")
        rows.extend(
            [
                f"   * - ``0x{command}``",
                f"     - :doc:`{name} <{Path(filename).stem}>`",
                f"     - {(status or '未标注').split('。', 1)[0]}",
                f"     - {source}",
            ]
        )
    return "\n".join(
        [
            heading("3 L2 Command 详解"),
            "",
            *rows,
            "",
            ".. note::",
            "",
            "   原 ``0x07`` 控制命令已并入 ``0x04`` 通知命令；``0x0B`` 是文件传输命令的原编号。两者均保留不用。"
            "``0x0E``\\ （视频流）与 ``0x0F``\\ （远程控制）已预留，规范待补。",
            "",
            *(
                [
                    ".. warning:: 工作簿引用缺口",
                    "",
                    "   源工作簿提及以下工作表，但当前文件未包含："
                    + "、".join(f"``{name}``" for name in sorted(missing_sheets))
                    + "。相关描述保留原文，不能据此推断缺失的格式。",
                    "",
                ]
                if missing_sheets else []
            ),
            # Outside any subsection, so commands are second-level sidebar entries.
            ".. toctree::",
            "   :maxdepth: 1",
            "   :hidden:",
            "",
            *[f"   {Path(filename).stem}" for _, _, filename, _, _ in commands],
            "",
        ]
    )


def validate(wb) -> tuple[dict, list[str]]:
    examples = 0
    errors = []
    formulas = 0
    formula_errors = []
    missing_references: dict[str, list[str]] = {}
    sheet_names = set(wb.sheetnames) | {"语言编码"}
    command_by_sheet = {
        document.sheet: document.command
        for document in DOCUMENTS
        if document.command is not None
    }

    reference_patterns = (
        re.compile(r'见\s*[“"]([^”"]+)[”"]\s*工作表'),
        re.compile(r"(ws-[A-Za-z0-9_-]+\.md)"),
    )
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    normalized = " ".join(value.split())
                    expected_command = command_by_sheet.get(ws.title)
                    if PACKET.fullmatch(normalized) and expected_command:
                        tokens = HEX_PAIR.findall(normalized)
                        if tokens[0].upper() != expected_command.upper():
                            continue
                        examples += 1
                        declared = int(tokens[3] + tokens[4], 16)
                        actual = len(tokens) - 5
                        if declared != actual:
                            errors.append(
                                f"{ws.title}!{cell.coordinate}: declared v-length {declared}, actual {actual}"
                            )
                    for pattern in reference_patterns:
                        for reference in pattern.findall(value):
                            if reference.endswith(".md") or reference in sheet_names:
                                continue
                            missing_references.setdefault(reference, []).append(
                                f"{ws.title}!{cell.coordinate}"
                            )
                    if value.startswith("="):
                        formulas += 1
                    if any(marker in value for marker in ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?")):
                        formula_errors.append(f"{ws.title}!{cell.coordinate}")

    report = {
        "complete_command_examples": examples,
        "formula_count": formulas,
        "formula_errors": formula_errors,
        "missing_referenced_sheets": missing_references,
    }
    return report, errors


def build_outputs() -> tuple[dict[Path, str], dict]:
    if not SOURCE.exists():
        raise FileNotFoundError(f"source workbook not found: {SOURCE}")
    wb = load_workbook(SOURCE, data_only=False)
    missing_sheets = [document.sheet for document in DOCUMENTS if document.sheet not in wb.sheetnames]
    if missing_sheets:
        raise ValueError(f"mapped sheets not found: {', '.join(missing_sheets)}")

    outputs: dict[Path, str] = {}
    summaries = []
    for document in DOCUMENTS:
        content, status = render_sheet(document, wb[document.sheet])
        path = OUTPUT_DIR / document.filename
        outputs[path] = content
        summaries.append(
            {
                "sheet": document.sheet,
                "filename": document.filename,
                "title": document.title,
                "command": document.command,
                "status": status,
                "rows": wb[document.sheet].max_row,
            }
        )
    report, errors = validate(wb)
    outputs[OUTPUT_DIR / "index.rst"] = render_index(summaries, report["missing_referenced_sheets"])
    manifest = {
        "source": str(SOURCE.relative_to(ROOT)),
        "source_sha256": sha256(SOURCE),
        "documents": summaries,
        "checks": {
            **report,
            "vlength_errors": errors,
        },
    }
    outputs[MANIFEST] = json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    if errors:
        raise ValueError("; ".join(errors))
    return outputs, manifest


def compare(outputs: dict[Path, str]) -> list[str]:
    stale = []
    obsolete = [OUTPUT_DIR / "index.md"] + [
        OUTPUT_DIR / Path(document.filename).with_suffix(".md").name
        for document in DOCUMENTS
    ]
    stale.extend(
        str(path.relative_to(ROOT)) for path in obsolete if path.exists()
    )
    for path, content in outputs.items():
        if not path.exists() or path.read_text(encoding="utf-8") != content:
            stale.append(str(path.relative_to(ROOT)))
    return stale


def write(outputs: dict[Path, str]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for path, content in outputs.items():
        path.write_text(content, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="fail if generated pages are stale")
    args = parser.parse_args()

    outputs, manifest = build_outputs()
    if args.check:
        stale = compare(outputs)
        if stale:
            print("generated protocol pages are stale:")
            for path in stale:
                print(f"  {path}")
            return 2
        print("protocol pages are up to date")
    else:
        write(outputs)
        print(f"generated {len(outputs)} protocol files from {SOURCE.relative_to(ROOT)}")

    checks = manifest["checks"]
    print(
        f"checked {checks['complete_command_examples']} complete command examples; "
        f"v-length errors: {len(checks['vlength_errors'])}"
    )
    missing = checks["missing_referenced_sheets"]
    if missing:
        print("warning: workbook references sheets not included in the source:")
        for name, locations in sorted(missing.items()):
            print(f"  {name}: {', '.join(locations[:6])}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (FileNotFoundError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        sys.exit(1)
