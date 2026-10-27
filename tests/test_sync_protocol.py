import unittest

from openpyxl import load_workbook

from tools.sync_protocol import (
    DOCUMENTS,
    MANUAL_DOCUMENTS,
    OUTPUT_DIR,
    ROOT,
    RST_SOURCED_DOCUMENTS,
    SOURCE,
    build_outputs,
    heading,
)


MANUAL_COMMANDS = (
    "0x01-firmware-upgrade.rst",
    "0x06-factory-test.rst",
    "0x0c-ble-conn-param.rst",
    "0x0d-wifi-provision.rst",
    "0x10-file-transfer.rst",
)

LEGACY_PROTOCOL_MARKERS = ("旧版协议", "旧协议", "老协议", "原编号", "原来的编号", "原定")


class ProtocolRenderingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.outputs, cls.manifest = build_outputs()

    @classmethod
    def page(cls, name):
        path = OUTPUT_DIR / name
        if path in cls.outputs:
            return cls.outputs[path]
        return path.read_text(encoding="utf-8")

    def test_opening_sections_are_not_dropped(self):
        for name in ("0x02-settings.rst", "0x03-bind.rst"):
            page = self.page(name)
            self.assertIn("命令格式\n--------\n", page)
            self.assertLess(page.index("命令格式\n--------"), page.index(".. list-table::"))

    def test_generated_chapters_have_one_navigation_level(self):
        for name in (
            "0x02-settings.rst",
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            with self.subTest(name=name):
                page = self.page(name)
                self.assertTrue(page.startswith(":tocdepth: 2\n\n"))
                self.assertRegex(page, r"\n命令格式\n-+\n")
                self.assertRegex(page, r"\n具体命令\n-+\n")
                self.assertNotRegex(page, r"(?m)^1\.\d+ .+\n~+")
                self.assertNotRegex(page, r"(?m)^\d+\. 通用说明\n-+")
        for name in (
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            self.assertRegex(self.page(name), r"\n规则\n-+\n")
        self.assertIn("key-interaction", self.page("0x04-notification.rst"))
        self.assertIn("key-interaction", self.page("0x11-measurement.rst"))

    def test_generated_chapters_do_not_show_source_metadata(self):
        for document in DOCUMENTS:
            with self.subTest(name=document.filename):
                page = self.page(document.filename)
                self.assertNotIn("container:: doc-meta", page)
                self.assertNotIn("工作表（自动同步）", page)

    def test_rst_sourced_command_index_lists_rst_origin(self):
        index = self.page("index.rst")
        self.assertIn("``0x05``", index)
        self.assertIn("第一版本", index)
        self.assertIn("RST", index)
        self.assertIn("0x05-data.rst", self.manifest["rst_sourced"])

    def test_manual_key_details_are_nested_under_chapters(self):
        for name in MANUAL_COMMANDS:
            with self.subTest(name=name):
                page = (OUTPUT_DIR / name).read_text(encoding="utf-8")
                self.assertTrue(page.startswith(":tocdepth: 2\n\n"))
                self.assertRegex(page, r"\n规则\n-+\n")
                self.assertRegex(page, r"\n具体命令\n-+\n")
                self.assertNotRegex(page, r"(?m)^0x[\da-fA-F/]+ — .+\n-+")
        self.assertIn(".. table:: 消息定义", (OUTPUT_DIR / "0x08-dump-stack.rst").read_text())

    def test_firmware_upgrade_starts_with_command_format(self):
        page = (OUTPUT_DIR / "0x01-firmware-upgrade.rst").read_text(encoding="utf-8")
        self.assertIn("命令格式\n--------\n", page)
        self.assertIn(".. list-table:: L2 命令格式\n", page)
        command_format = page.index("命令格式\n--------")
        rules = page.index("规则\n----")
        concrete_commands = page.index("具体命令\n--------")
        self.assertLess(command_format, rules)
        self.assertLess(rules, concrete_commands)
        self.assertIn("   * - Length\n     - 2B\n     - Value 字节数\n     - Big-Endian", page)

    def test_all_command_pages_share_the_same_section_skeleton(self):
        expected = ["命令格式", "规则", "具体命令"]
        for path in sorted(OUTPUT_DIR.glob("0x*.rst")):
            lines = path.read_text(encoding="utf-8").splitlines()
            sections = []
            for index in range(len(lines) - 1):
                title = lines[index].strip()
                underline = lines[index + 1].strip()
                if (
                    title
                    and not title.startswith("*")
                    and len(underline) >= 3
                    and set(underline) == {"-"}
                ):
                    sections.append(title)
            with self.subTest(path=path.name):
                self.assertEqual(sections, expected)

    def test_concrete_command_sections_use_protocol_example_template(self):
        for path in sorted(OUTPUT_DIR.glob("0x*.rst")):
            text = path.read_text(encoding="utf-8")
            section = text.split("具体命令\n--------", 1)[1]
            expected_headers = (
                ("Key", "opcode", "名称 / 用途", "方向", "具体的命令", "数据说明")
                if path.name == "0x02-settings.rst"
                else ("Key", "名称 / 用途", "方向", "具体的命令", "数据说明")
            )
            with self.subTest(path=path.name):
                expected_widths = "5 10 17 9 15 44" if path.name == "0x02-settings.rst" else "5 20 10 15 50"
                expected_columns = "columns-6" if path.name == "0x02-settings.rst" else "columns-5"
                self.assertIn(f":widths: {expected_widths}", section)
                self.assertIn(f":class: protocol-table command-detail {expected_columns}", section)
                header_block = section.split("   * - ", 1)[1].split("\n   * - ", 1)[0]
                header_values = [
                    line.removeprefix("     - ") if line.startswith("     - ") else line.strip()
                    for line in header_block.splitlines()
                ]
                self.assertEqual(tuple(header_values), expected_headers)

    def test_command_detail_tables_reserve_space_for_key_column(self):
        stylesheet = (ROOT / "docs/_static/custom.css").read_text(encoding="utf-8")
        self.assertIn(
            ".rst-content table.protocol-table.command-detail td:nth-child(1) {\n"
            "  overflow-wrap: normal;\n"
            "  padding-right: 0.5rem;\n"
            "  padding-left: 0.5rem;\n"
            "  white-space: normal;",
            stylesheet,
        )
        self.assertIn(
            ".columns-5 col:first-child {\n  width: 10% !important;",
            stylesheet,
        )
        self.assertIn(
            ".columns-5 col:nth-child(2) {\n  width: 15% !important;",
            stylesheet,
        )
        self.assertIn(
            ".columns-6 col:first-child {\n  width: 7% !important;",
            stylesheet,
        )
        self.assertIn(
            ".columns-6 col:nth-child(3) {\n  width: 15% !important;",
            stylesheet,
        )
        self.assertIn(
            "@media screen and (max-width: 1279px) {\n"
            "  .rst-content .wy-table-responsive,\n"
            "  .rst-content .table-wrapper {\n"
            "    overflow-x: auto;",
            stylesheet,
        )
        self.assertIn(".columns-5 {\n    min-width: 760px;", stylesheet)
        self.assertIn(".columns-6 {\n    min-width: 900px;", stylesheet)

    def test_command_detail_tables_use_one_key_per_row(self):
        for path in sorted(OUTPUT_DIR.glob("0x*.rst")):
            section = path.read_text(encoding="utf-8").split("具体命令\n--------", 1)[1]
            with self.subTest(path=path.name):
                self.assertNotRegex(section, r"(?m)^   \* - 0x[\da-fA-F]{2}\s*/")

    def test_manual_rule_sections_do_not_repeat_command_catalogs(self):
        for name in (
            "0x01-firmware-upgrade.rst",
            "0x06-factory-test.rst",
            "0x0c-ble-conn-param.rst",
            "0x0d-wifi-provision.rst",
            "0x10-file-transfer.rst",
        ):
            text = (OUTPUT_DIR / name).read_text(encoding="utf-8")
            rules = text.split("规则\n----", 1)[1].split("具体命令\n--------", 1)[0]
            with self.subTest(name=name):
                self.assertNotIn("消息定义", rules)
                self.assertNotRegex(rules, r"(?m)^0x[\da-fA-F]+ — .+\n~+")

    def test_document_body_lists_use_ordered_numbers(self):
        for path in (ROOT / "docs").rglob("*.rst"):
            content = path.read_text(encoding="utf-8")
            with self.subTest(path=str(path.relative_to(ROOT))):
                self.assertNotRegex(content, r"(?m)^\*\s+")
                self.assertNotRegex(content, r"(?m)^-\s+")

    def test_tocdepth_precedes_title_on_every_command_page(self):
        generated = {document.filename for document in DOCUMENTS if document.command}
        filenames = sorted(generated)
        filenames.extend(filename for filename, _, _ in MANUAL_DOCUMENTS)
        for name in filenames:
            with self.subTest(name=name):
                page = (
                    self.page(name)
                    if name in generated
                    else (OUTPUT_DIR / name).read_text(encoding="utf-8")
                )
                self.assertTrue(page.startswith(":tocdepth: 2\n\n0x"), name)

    def test_command_pages_do_not_repeat_l2_version(self):
        generated = {document.filename for document in DOCUMENTS if document.command}
        filenames = sorted(generated)
        filenames.extend(filename for filename, _, _ in MANUAL_DOCUMENTS)
        for name in filenames:
            with self.subTest(name=name):
                page = (
                    self.page(name)
                    if name in generated
                    else (OUTPUT_DIR / name).read_text(encoding="utf-8")
                )
                self.assertNotIn("L2 版本号", page)

    def test_command_field_uses_length_name(self):
        for path in sorted(OUTPUT_DIR.glob("0x*.rst")):
            with self.subTest(path=path.name):
                self.assertNotIn("v-length", path.read_text(encoding="utf-8"))
        for name in ("architecture.rst", "conventions.rst"):
            with self.subTest(name=name):
                page = (ROOT / "docs" / name).read_text(encoding="utf-8")
                self.assertNotIn("v-length", page.lower())

    def test_internal_section_references_follow_renamed_headings(self):
        data = self.page("0x05-data.rst")
        self.assertIn("见 :ref:`具体命令 <cmd-05-section-3>`", data)
        bind = self.page("0x03-bind.rst")
        self.assertIn("见 :ref:`规则 <cmd-03-section-2>`", bind)
        self.assertNotIn('见"二、消息定义"', bind)
        settings = self.page("0x02-settings.rst")
        self.assertIn("见 :ref:`规则 <cmd-02-section-2>`", settings)
        for name in ("0x04-notification.rst", "0x05-data.rst", "0x11-measurement.rst"):
            self.assertNotIn("见:ref:", self.page(name))
        # Every section anchor must be referenced with a :ref: rather than a
        # hard-coded section number, which the toctree numbering does not honour.
        for name in (
            "0x02-settings.rst",
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            with self.subTest(name=name):
                self.assertNotRegex(self.page(name), r"见第\s*\d+(?:～\d+)?\s*节")

    def test_packet_examples_are_tables_even_when_source_lines_join(self):
        for name in (
            "0x02-settings.rst",
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            with self.subTest(name=name):
                self.assertIn(".. list-table:: 协议数据示例", self.page(name))
        bind = self.page("0x03-bind.rst")
        self.assertIn("   * - 0x02\n     - 绑定返回 / 绑定结果", bind)
        self.assertNotIn("Value03 00", bind)
        self.assertNotIn("操作数据02 00", self.page("0x02-settings.rst"))

    def test_appendices_include_headers_and_first_entries(self):
        weather = self.page("weather-types.rst")
        self.assertTrue(weather.startswith(":orphan:\n\n"))
        self.assertIn("   * - 取值\n     - 天气\n     - 英文", weather)
        self.assertIn("   * - 0x00\n     - 其他\n     - Other", weather)
        language = self.page("language-codes.rst")
        self.assertTrue(language.startswith(":orphan:\n\n"))
        self.assertIn("   * - 取值\n     - 语言\n     - 英文", language)
        self.assertIn("   * - 0x00（0）\n     - 英文\n     - English", language)

    def test_appendix_is_a_separate_chapter_with_workbook_download(self):
        root_index = (ROOT / "docs/index.rst").read_text(encoding="utf-8")
        appendix = (ROOT / "docs/appendix.rst").read_text(encoding="utf-8")
        self.assertIn("   protocol/index\n   appendix\n", root_index)
        self.assertTrue(appendix.startswith("附表\n"))
        self.assertIn(
            f":download:`协议工作簿（Excel） <assets/protocol/{SOURCE.name}>`",
            appendix,
        )
        self.assertTrue(SOURCE.is_file())
        command_index = self.page("index.rst")
        self.assertNotIn("附表", command_index)
        self.assertNotIn("   weather-types", command_index)
        self.assertNotIn("   language-codes", command_index)

    def test_chapter_and_section_numbers_are_generated_by_toctree(self):
        root_index = (ROOT / "docs/index.rst").read_text(encoding="utf-8")
        self.assertIn("   :includehidden:\n   :numbered:\n", root_index)
        for filename, title in (
            ("conventions.rst", "名词解释与约定"),
            ("architecture.rst", "协议结构介绍"),
            ("appendix.rst", "附表"),
        ):
            source = (ROOT / "docs" / filename).read_text(encoding="utf-8")
            self.assertTrue(source.startswith(f"{title}\n"), filename)
        self.assertTrue(self.page("index.rst").startswith("L2 Command 详解\n"))

    def test_intro_chapters_use_concise_unnumbered_section_titles(self):
        chapters = {
            "conventions.rst": ("名词解释", "约定", "CRC16 计算范围"),
            "architecture.rst": ("协议栈结构", "L0 接入承载", "L1 传输层", "L2 应用层"),
        }
        for filename, titles in chapters.items():
            page = (ROOT / "docs" / filename).read_text(encoding="utf-8")
            with self.subTest(filename=filename):
                for title in titles:
                    self.assertIn(heading(title, 2), page)
                self.assertNotRegex(page, r"(?m)^[12]\.\d+(?:\.\d+)? .+\n[-~]+")
        architecture = (ROOT / "docs/architecture.rst").read_text(encoding="utf-8")
        self.assertIn(heading("BLE UART Profile", 3), architecture)
        self.assertNotIn(heading("GATT 服务 UUID", 3), architecture)
        for anchor in ("l0-uart-profile-ble", "l1-transport-layer", "l2-application-layer"):
            self.assertIn(f".. _{anchor}:", architecture)

    def test_architecture_does_not_impose_a_20_byte_mtu_limit(self):
        architecture = (ROOT / "docs/architecture.rst").read_text(encoding="utf-8")
        self.assertNotIn("ATT MTU", architecture)
        self.assertNotIn("20 字节", architecture)
        self.assertNotIn("20 bytes", architecture)
        self.assertIn("L1 通过拆包/组包支持最大 **65535 bytes** payload", architecture)

    def test_l2_payload_table_uses_compact_unit_layout(self):
        architecture = (ROOT / "docs/architecture.rst").read_text(encoding="utf-8")
        self.assertIn("L2 Payload 由连续的 Key-Value 单元组成", architecture)
        self.assertIn(".. list-table:: L2 Payload 单元结构", architecture)
        self.assertIn(":class: packet-table payload-unit-table", architecture)
        self.assertNotIn("* - Key-Value 1", architecture)

    def test_short_bitfield_tables_use_centered_cells(self):
        architecture = (ROOT / "docs/architecture.rst").read_text(encoding="utf-8")
        self.assertEqual(architecture.count(":class: packet-table bitfield-table"), 2)
        self.assertIn("   * - Command ID\n     - 8 bits\n     - 命令类型，见 Command 列表", architecture)
        self.assertIn("   * - Length\n     - 16 bits\n     - 该 key 的 value 字节数，最大 65535", architecture)

    def test_table_names_are_captions_instead_of_section_headings(self):
        architecture = (ROOT / "docs" / "architecture.rst").read_text(encoding="utf-8")
        for heading_text, caption in (
            ("L1 Header（8 bytes，Big-Endian）", ".. table:: L1 Header（8 bytes，Big-Endian）"),
            ("ACK 包格式", ".. table:: ACK 包格式"),
            ("L2 Header（16 bits，Big-Endian）", ".. list-table:: L2 Header（16 bits，Big-Endian）"),
            ("Key Header（16 bits）", ".. list-table:: Key Header（16 bits）"),
        ):
            with self.subTest(heading=heading_text):
                self.assertNotIn(heading(heading_text, 3), architecture)
                self.assertIn(caption, architecture)
        self.assertIn(".. table:: GATT 服务 UUID", architecture)

        wifi = (OUTPUT_DIR / "0x0d-wifi-provision.rst").read_text(encoding="utf-8")
        self.assertIn(".. list-table:: error 错误码\n", wifi)
        self.assertIn(".. list-table:: state 状态码\n", wifi)
        self.assertIn(".. list-table:: error 错误码（state=0x03）\n", wifi)

        for name, caption in (
            ("0x02-settings.rst", ".. list-table:: 操作码定义\n"),
            ("0x03-bind.rst", ".. list-table:: 协议数据示例\n"),
            ("0x04-notification.rst", ".. list-table:: Key 分组\n"),
            ("0x0a-log.rst", ".. list-table:: L2 命令格式\n"),
            ("0x11-measurement.rst", ".. list-table:: Key 分组\n"),
            ("0x12-simulator.rst", ".. list-table:: 生成规则（按类型）\n"),
            ("0x13-audio.rst", ".. list-table:: 协议数据示例\n"),
        ):
            with self.subTest(name=name):
                self.assertIn(caption, self.page(name))

    def test_table_captions_use_numbered_upright_style(self):
        config = (ROOT / "docs/conf.py").read_text(encoding="utf-8")
        self.assertIn("numfig = True", config)
        self.assertIn("numfig_secnum_depth = 2", config)
        self.assertIn('numfig_format = {"table": "表 %s"}', config)
        stylesheet = (ROOT / "docs/_static/custom.css").read_text(encoding="utf-8")
        self.assertIn(".rst-content table.docutils > caption {", stylesheet)
        self.assertIn("text-align: center;", stylesheet)
        self.assertIn("font-style: normal;", stylesheet)
        self.assertNotIn("font-style: italic;", stylesheet)

    def test_every_documented_table_has_a_caption(self):
        documents = {path: content for path, content in self.outputs.items() if path.suffix == ".rst"}
        documents.update(
            {
                path: path.read_text(encoding="utf-8")
                for path in (ROOT / "docs").rglob("*.rst")
            }
        )
        for path, content in documents.items():
            for line_number, line in enumerate(content.splitlines(), start=1):
                if not line.startswith((".. list-table::", ".. table::")):
                    continue
                with self.subTest(path=str(path.relative_to(ROOT)), line=line_number):
                    self.assertRegex(line, r"^\.\. (?:list-table|table)::\s+\S")

    def test_redundant_table_explanations_are_removed(self):
        phrases = (
            "以下为 L2 层命令格式",
            '"具体的命令"一列为完整的 L2 数据',
            "完整过程见\"交互示例\"工作表",
            "需要固件确认的问题见\"待确认\"工作表",
            "需要固件确认和待讨论的问题见\"待确认\"工作表",
        )
        documents = {path: content for path, content in self.outputs.items() if path.suffix == ".rst"}
        documents.update(
            {
                path: path.read_text(encoding="utf-8")
                for path in (ROOT / "docs").rglob("*.rst")
            }
        )
        for path, content in documents.items():
            for phrase in phrases:
                with self.subTest(path=str(path.relative_to(ROOT)), phrase=phrase):
                    self.assertNotIn(phrase, content)

        workbook = load_workbook(SOURCE, read_only=True, data_only=False)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for phrase in phrases:
                        with self.subTest(sheet=worksheet.title, cell=cell.coordinate, phrase=phrase):
                            self.assertNotIn(phrase, cell.value)

    def test_heading_hierarchy_is_visually_distinct(self):
        stylesheet = (ROOT / "docs/_static/custom.css").read_text(encoding="utf-8")
        self.assertIn(".rst-content h1 {\n  margin: 0 0 1.4rem;", stylesheet)
        self.assertIn(".rst-content h2 {\n  margin: 2.75rem 0 1.15rem;", stylesheet)
        self.assertIn("border-left: 4px solid #1f6f95;", stylesheet)
        self.assertIn("font-size: 1.48rem;", stylesheet)
        self.assertIn(".rst-content h3 {\n  margin: 1.9rem 0 0.8rem;", stylesheet)
        self.assertIn("font-size: 1.04rem;", stylesheet)

    def test_body_typography_prioritizes_chinese_reading(self):
        stylesheet = (ROOT / "docs/_static/custom.css").read_text(encoding="utf-8")
        self.assertIn('"Source Han Sans SC", "Noto Sans CJK SC", "Noto Sans SC"', stylesheet)
        self.assertIn("--doc-reading-width: 48rem;", stylesheet)
        self.assertIn(".rst-content {\n  color: var(--doc-text);\n  font-size: 17px;\n  line-height: 1.75;", stylesheet)
        self.assertIn(".rst-content p {\n  max-width: var(--doc-reading-width);", stylesheet)
        self.assertIn(
            ".rst-content .section ul,\n.rst-content .section ol {\n  max-width: var(--doc-reading-width);",
            stylesheet,
        )
        self.assertIn("margin-bottom: 1.25rem;\n  line-height: 1.75;", stylesheet)
        self.assertIn("margin-top: 0.35rem;\n  margin-bottom: 1.5rem;", stylesheet)
        self.assertIn("margin-bottom: 0.5rem;\n  line-height: 1.75;", stylesheet)
        self.assertIn("font-size: 14px;", stylesheet)
        self.assertIn(".rst-content {\n    font-size: 16px;\n    line-height: 1.72;", stylesheet)
        self.assertIn(".rst-content p,\n  .rst-content li {\n    line-height: 1.72;", stylesheet)

    def test_architecture_stack_uses_structured_protocol_map(self):
        architecture = (ROOT / "docs/architecture.rst").read_text(encoding="utf-8")
        self.assertIn('<figure class="protocol-stack"', architecture)
        self.assertNotIn("block-beta", architecture)
        for layer in ("L2", "L1", "L0"):
            with self.subTest(layer=layer):
                self.assertIn(f'<span class="protocol-layer__id">{layer}</span>', architecture)
        for adapter in ("ble", "uart", "socket"):
            with self.subTest(adapter=adapter):
                self.assertIn(f"protocol-adapter--{adapter}", architecture)
        for carrier in ("BLE", "UART", "SOCKET"):
            with self.subTest(carrier=carrier):
                self.assertIn(f">{carrier}</span>", architecture)

    def test_flow_diagrams_and_references_render(self):
        data = self.page("0x05-data.rst")
        self.assertIn(":doc:`点测命令 <0x11-measurement>`", data)
        self.assertIn("同步全部", data)
        self.assertIn(":doc:`天气类型 <weather-types>`", self.page("0x02-settings.rst"))
        for name in ("0x04-notification.rst", "0x05-data.rst", "0x11-measurement.rst"):
            self.assertNotIn("ws-protocol-", self.page(name))
        # A role wrapped in literal markup renders verbatim instead of linking.
        for name in (
            "0x02-settings.rst",
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            with self.subTest(name=name):
                self.assertNotIn("``:doc:", self.page(name))

    def test_protocol_docs_do_not_compare_against_legacy_protocols(self):
        document_paths = set(self.outputs) | set(OUTPUT_DIR.glob("*.rst"))
        document_paths.update(
            path
            for path in (ROOT / "docs").glob("*.rst")
            if path.name in {"architecture.rst", "conventions.rst", "index.rst"}
        )
        for path in document_paths:
            if path.suffix not in {".rst", ".json"}:
                continue
            content = path.read_text(encoding="utf-8")
            for marker in LEGACY_PROTOCOL_MARKERS:
                with self.subTest(path=str(path.relative_to(ROOT)), marker=marker):
                    self.assertNotIn(marker, content)

        workbook = load_workbook(SOURCE, read_only=True, data_only=False)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for marker in LEGACY_PROTOCOL_MARKERS:
                        with self.subTest(sheet=worksheet.title, cell=cell.coordinate, marker=marker):
                            self.assertNotIn(marker, cell.value)

    def test_rst_sourced_pages_no_longer_reference_workbook_sheets(self):
        # 0x02～0x05、0x0A、0x11～0x13 now come from HoneyBox markdown, so the
        # workbook no longer has to carry the sheets those pages used to cite.
        missing = self.manifest["checks"]["missing_referenced_sheets"]
        self.assertEqual(missing, {})
        self.assertEqual(self.manifest["checks"]["formula_errors"], [])
        index = self.page("index.rst")
        self.assertNotIn(".. warning::", index)
        self.assertEqual(
            self.manifest["rst_sourced"],
            [
                "0x02-settings.rst",
                "0x03-bind.rst",
                "0x04-notification.rst",
                "0x05-data.rst",
                "0x0a-log.rst",
                "0x11-measurement.rst",
                "0x12-simulator.rst",
                "0x13-audio.rst",
            ],
        )

    def test_admonitions_do_not_repeat_manual_titles(self):
        document_paths = set(self.outputs) | set(OUTPUT_DIR.glob("*.rst"))
        document_paths.update((ROOT / "docs").glob("*.rst"))
        for path in document_paths:
            if path.suffix != ".rst":
                continue
            content = path.read_text(encoding="utf-8")
            for line in content.splitlines():
                with self.subTest(path=str(path.relative_to(ROOT)), line=line):
                    self.assertFalse(line.startswith(".. note:: ") or line.startswith(".. warning:: "))

    def test_heading_underline_covers_wide_characters(self):
        title, underline = heading("二、消息定义", 2).splitlines()
        self.assertEqual(title, "二、消息定义")
        self.assertGreater(len(underline), len(title))

    def test_command_format_tables_use_readable_column_widths(self):
        for name in (
            "0x02-settings.rst",
            "0x03-bind.rst",
            "0x04-notification.rst",
            "0x05-data.rst",
            "0x0a-log.rst",
            "0x11-measurement.rst",
            "0x12-simulator.rst",
            "0x13-audio.rst",
        ):
            with self.subTest(name=name):
                page = self.page(name)
                self.assertIn(":widths: 18 10 18 54", page)
                self.assertIn(
                    ":class: protocol-table command-format columns-4", page
                )


if __name__ == "__main__":
    unittest.main()
